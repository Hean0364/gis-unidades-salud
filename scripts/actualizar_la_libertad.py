"""Completa el sector de La Libertad y aplica el polígono aportado para San Mateo.

No ejecuta pruebas ni modifica los otros Excel departamentales.
"""

from collections import Counter
from copy import copy
import unicodedata

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn

from generar_mapas_departamentales import DEPARTMENTS, collection, geometry, records, write_layer
from preparar_fuentes import geom_centroid


SAN_MATEO_WKT = "POLYGON ((-89.36839387839004 13.72482961806833, -89.36854456317657 13.72453659473763, -89.36844586330508 13.72448125127536, -89.36828879937505 13.72478057309364, -89.36839387839004 13.72482961806833))"
PRIVATE_SOURCES = {
    "way/378234332": "https://www.facebook.com/p/Cl%C3%ADnica-Corrales-100083695633545/",
    "way/397372547": "https://www.openstreetmap.org/way/397372547",
    "way/397493073": "https://www.openstreetmap.org/way/397493073",
    "way/398033509": "https://hlourdes.com/",
    "way/398052521": "https://www.facebook.com/p/Laboratorio-Cl%C3%ADnico-Clio-Med-61565687895844/",
    "way/398568020": "https://www.labcliza.com/",
    "way/398606707": "https://www.facebook.com/Jawetz.Laboratorio/",
    "way/399286333": "https://www.facebook.com/clinicamedicafamiliar1/",
    "way/399727751": "https://www.facebook.com/hospitalmerliot/",
    "way/442217631": "https://www.hospitaldiagnostico.com/centro-especialidades/6",
    "way/541308878": "https://policlinicacasadesalud.com/",
    "way/693779698": "https://www.ordendemaltaelsalvador.org/",
    "way/1110828563": "https://www.fundesoelsalvador.org/",
    "way/1152815588": "https://www.fundesoelsalvador.org/",
    "way/1200045596": "https://www.facebook.com/p/clinica-asistencial-medica-cooperativa-Chanmico-61585440324158/",
    "way/1215703670": "https://www.facebook.com/Biogenesis.sv/",
    "way/1320205287": "https://ecas.edu.sv/",
    "way/1382672260": "https://www.hospitalclimosal.com/",
    "node/661971230": "https://hsanmateo.com/",
    "node/4003452952": "https://www.facebook.com/TratamientosdelPieCAIP/",
    "node/4001465053": "https://www.facebook.com/clinicksusalud/",
}


def normalize(value):
    return "".join(char for char in unicodedata.normalize("NFD", str(value or "").lower()) if not unicodedata.combining(char))


def classify(record):
    name = normalize(record.get("nombre"))
    source_id = record.get("fuente_id")
    if record.get("fuente") == "manual" and record.get("sector") in {"Público", "Privado"} and record.get("fuente_sector"):
        return record["sector"], record.get("criterio_sector") or "Clasificación manual documentada.", record["fuente_sector"]
    if record.get("fuente") == "minsal" or "unidad de salud" in name or "unidad de salut" in name or "unidad comunitaria salud familiar" in name:
        return "Público", "Establecimiento identificado de la red pública MINSAL/UCSF.", "https://www.salud.gob.sv/"
    if "hospital nacional san rafael" in name or "anexo hospital san rafael" in name:
        return "Público", "Hospital Nacional San Rafael y sus instalaciones anexas: red pública MINSAL.", "https://www.salud.gob.sv/"
    if "isss" in name or "especializado integral de atencion ambulatoria la ceiba" in name or name in {"clinica comunal santa monica", "clinica comununal de antiguo cuscatlan", "clinica comunal de antiguo cuscatlan"}:
        return "Público", "Centro propio o clínica comunal del Instituto Salvadoreño del Seguro Social.", "https://www.isss.gob.sv/centros-de-atencion/"
    if "isbm" in name or "policlinico magisterial" in name:
        return "Público", "Centro identificado del Instituto Salvadoreño de Bienestar Magisterial.", "https://www.isbm.gob.sv/policlinicos/"
    if "clinica municipal" in name:
        return "Público", "Clínica identificada como municipal, de administración local.", record.get("fuente_url")
    if any(term in name for term in ["cruz roja", "cruz verde", "comandos de salvamento"]):
        return "Privado", "Organización humanitaria no estatal, sin fines de lucro; prestar auxilio público no la convierte en entidad estatal.", record.get("fuente_url")
    if source_id in PRIVATE_SOURCES:
        nonprofit = any(term in name for term in ["profamilia", "parroquial", "orden de malta", "fundeso", "ecas", "chanmico"])
        criterion = "Entidad no estatal de carácter social, religioso, fundacional o cooperativo; no se clasifica como pública por ofrecer atención social." if nonprofit else "Prestador médico o laboratorio no estatal identificado por nombre y ubicación en su publicación institucional."
        return "Privado", criterion, PRIVATE_SOURCES[source_id]
    return "No determinado", "La información disponible no permite identificar con certeza la titularidad. No se supone que sea privado por llamarse clínica u hospital, ni por carecer de operador.", record.get("fuente_url")


def main():
    folder = DEPARTMENTS / "La Libertad"
    path = folder / "establecimientos_salud.xlsx"
    workbook = load_workbook(path)
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    old_max_column = sheet.max_column
    for name in ["sector", "criterio_sector", "fuente_sector", "operador_original"]:
        if name not in columns:
            column = sheet.max_column + 1
            columns[name] = column
            cell = sheet.cell(1, column, name)
            cell._style = copy(sheet.cell(1, columns["operador"])._style)
            sheet.column_dimensions[get_column_letter(column)].width = 22 if name in {"sector", "operador_original"} else 70
    counts = Counter()
    for row_number in range(2, sheet.max_row + 1):
        record = {name: sheet.cell(row_number, column).value for name, column in columns.items()}
        if not record.get("id"):
            continue
        sector, criterion, reference = classify(record)
        for name, value in {"sector": sector, "criterio_sector": criterion, "fuente_sector": reference}.items():
            cell = sheet.cell(row_number, columns[name], value)
            cell._style = copy(sheet.cell(row_number, columns["operador"])._style)
        counts[sector] += 1
        # Corregir las etiquetas de operador público contradichas por la identidad
        # verificada, sin perder su valor original de las fuentes.
        if record.get("fuente_id") in {"way/398033509", "node/661971230"} and sector == "Privado":
            if record.get("operador") and record["operador"] != "Privado":
                if not record.get("operador_original"):
                    sheet.cell(row_number, columns["operador_original"], record["operador"])
                sheet.cell(row_number, columns["operador"], "Privado")
        if record["id"] == "hot-node-661971230":
            shape = geometry(SAN_MATEO_WKT)
            longitude, latitude = geom_centroid(shape)
            latitude, longitude = round(latitude, 10), round(longitude, 10)
            updates = {
                "geometria_tipo": "Polygon", "detalle": "Polígono", "detalle_codigo": "poligono",
                "vertices": len(shape["coordinates"][0]), "geometria_wkt": SAN_MATEO_WKT,
                "centroide_latitud": latitude, "centroide_longitud": longitude,
                "centroide_metodo": "centroide_area",
                "google_maps": f"https://www.google.com/maps/search/?api=1&query={latitude},{longitude}",
            }
            for name, value in updates.items():
                sheet.cell(row_number, columns[name], value)
            cell = sheet.cell(row_number, columns["geometria_wkt"])
            if cell.comment is None:
                cell.comment = Comment("Polígono suministrado por el usuario. Se conserva el ID del registro HOT original para trazabilidad; no se atribuye esta geometría al nodo OSM.", "Actualización de datos")
    for table in sheet.tables.values():
        if len(table.tableColumns) == old_max_column:
            for column in range(old_max_column + 1, sheet.max_column + 1):
                table.tableColumns.append(TableColumn(id=column, name=sheet.cell(1, column).value))
            table.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    workbook.save(path)
    workbook.close()
    data = collection("La Libertad", records(path))
    write_layer(folder / "establecimientos", "ESTABLECIMIENTOS_DEPARTAMENTO_GEOJSON", data)
    print(f"La Libertad: {dict(counts)}")
    print(f"Mapa individual: {data['meta']['count']} registros; {data['meta']['puntos']} puntos; {data['meta']['poligonos']} polígonos")
    print("Hospital San Mateo actualizado con el polígono exacto proporcionado.")


if __name__ == "__main__":
    main()
