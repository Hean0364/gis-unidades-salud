"""Crea el XLSX depurado de edificios internos y puntos sobre polígonos."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
INPUT_XLSX = DATA / "unidades_salud_limpias_excel" / "unidades_salud_sin_privados_con_fines_de_lucro.xlsx"
INPUT_GEOJSON = DATA / "mapa_san_salvador" / "unidades_salud_sin_privados_con_fines_de_lucro.geojson"
OUTPUT_XLSX = DATA / "unidades_salud_limpias_excel" / "unidades_salud_sin_privados_sin_superposiciones.xlsx"

BUILDING_IDS = {
    "unificado-geofabrik-way-1208406964",
    "unificado-geofabrik-way-1463787670",
    "unificado-geofabrik-way-1534982706",
}


def polygons(geometry: dict[str, Any]) -> list[list[list[list[float]]]]:
    if geometry["type"] == "Polygon":
        return [geometry["coordinates"]]
    if geometry["type"] == "MultiPolygon":
        return geometry["coordinates"]
    return []


def point_on_segment(
    point: list[float], start: list[float], end: list[float], epsilon: float = 1e-10
) -> bool:
    x, y = point
    x1, y1 = start
    x2, y2 = end
    cross = (x - x1) * (y2 - y1) - (y - y1) * (x2 - x1)
    if abs(cross) > epsilon:
        return False
    return (
        min(x1, x2) - epsilon <= x <= max(x1, x2) + epsilon
        and min(y1, y2) - epsilon <= y <= max(y1, y2) + epsilon
    )


def point_in_ring(point: list[float], ring: list[list[float]]) -> int:
    """Devuelve 0 fuera, 1 dentro y 2 sobre el borde."""
    x, y = point
    inside = False
    for start, end in zip(ring, ring[1:] + ring[:1]):
        if point_on_segment(point, start, end):
            return 2
        x1, y1 = start
        x2, y2 = end
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            inside = not inside
    return 1 if inside else 0


def polygon_covers_point(geometry: dict[str, Any], point: list[float]) -> bool:
    """Incluye interior y borde, y excluye el interior de los huecos."""
    for polygon in polygons(geometry):
        outer = point_in_ring(point, polygon[0])
        if outer == 2:
            return True
        if outer == 0:
            continue
        inside_hole = False
        for hole in polygon[1:]:
            hole_position = point_in_ring(point, hole)
            if hole_position == 2:
                return True
            if hole_position == 1:
                inside_hole = True
                break
        if not inside_hole:
            return True
    return False


def point_coordinates(geometry: dict[str, Any]) -> Iterable[list[float]]:
    if geometry["type"] == "Point":
        yield geometry["coordinates"]
    elif geometry["type"] == "MultiPoint":
        yield from geometry["coordinates"]


def exclusions() -> tuple[set[str], dict[str, list[str]], dict[str, str]]:
    collection = json.loads(INPUT_GEOJSON.read_text(encoding="utf-8"))
    features = collection["features"]
    features_by_id = {feature["properties"]["id"]: feature for feature in features}

    missing_buildings = BUILDING_IDS - set(features_by_id)
    if missing_buildings:
        raise ValueError(f"No se encontraron los tres edificios esperados: {sorted(missing_buildings)}")

    retained_polygons = [
        feature
        for feature in features
        if feature["geometry"]["type"] in {"Polygon", "MultiPolygon"}
        and feature["properties"]["id"] not in BUILDING_IDS
    ]
    point_hits: dict[str, list[str]] = {}
    names: dict[str, str] = {}

    for feature in features:
        geometry = feature["geometry"]
        if geometry["type"] not in {"Point", "MultiPoint"}:
            continue
        point_id = feature["properties"]["id"]
        names[point_id] = feature["properties"].get("nombre", "Sin nombre")
        containers = [
            polygon["properties"]["id"]
            for polygon in retained_polygons
            if any(
                polygon_covers_point(polygon["geometry"], coordinate)
                for coordinate in point_coordinates(geometry)
            )
        ]
        if containers:
            point_hits[point_id] = containers

    if len(BUILDING_IDS) != 3:
        raise ValueError("La depuración debe retirar exactamente tres edificios")
    if not point_hits:
        raise ValueError("No se detectaron puntos sobre polígonos")
    return set(point_hits), point_hits, names


def update_summary(workbook, retained: Counter[str], removed_points: int) -> None:
    summary = workbook["Resumen"]
    summary.delete_rows(1, summary.max_row)
    rows = [
        ["Unidades de salud sin privados lucrativos ni superposiciones", ""],
        ["Archivo de origen", str(INPUT_XLSX.relative_to(ROOT)).replace("/", "\\")],
        ["Registros conservados", sum(retained.values())],
        ["Edificios internos eliminados en esta actualización", len(BUILDING_IDS)],
        ["Puntos dentro o sobre polígonos eliminados", removed_points],
        [],
        ["Sector investigado", "Cantidad"],
        ["Pública", retained["Pública"]],
        ["Privada con fines de lucro", retained["Privada con fines de lucro"]],
        ["Privada sin fines de lucro", retained["Privada sin fines de lucro"]],
        ["No determinada", retained["No determinada"]],
        [],
        [
            "Filtros aplicados",
            "Se conservaron los registros del XLSX sin privados con fines de lucro; después se retiraron tres edificios internos y todos los puntos cubiertos por un polígono restante.",
        ],
        [
            "Regla espacial",
            "Un punto se retiró cuando su coordenada estaba dentro de un polígono o exactamente sobre su borde. Los tres edificios se retiraron antes de evaluar los puntos.",
        ],
    ]
    for row in rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 110
    summary.freeze_panes = "A2"
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True, size=14)
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    point_ids, point_hits, point_names = exclusions()
    workbook = load_workbook(INPUT_XLSX)
    removed = Counter()

    exclusions_by_sheet = {
        "Poligonos": BUILDING_IDS,
        "Puntos": point_ids,
    }
    for sheet_name, ids_to_remove in exclusions_by_sheet.items():
        worksheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        id_column = headers["id"]
        relation_column = headers.get("relacion_espacial_probable")

        for row_number in range(worksheet.max_row, 1, -1):
            record_id = worksheet.cell(row_number, id_column).value
            if record_id in ids_to_remove:
                worksheet.delete_rows(row_number)
                removed[sheet_name] += 1
            elif not record_id:
                worksheet.delete_rows(row_number)

        if sheet_name == "Poligonos" and relation_column:
            for row_number in range(2, worksheet.max_row + 1):
                relation = worksheet.cell(row_number, relation_column).value
                if relation and "hoja Puntos" in str(relation):
                    worksheet.cell(row_number, relation_column).value = None

        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        )

    if removed["Poligonos"] != 3:
        raise ValueError(f"Se eliminaron {removed['Poligonos']} edificios; se esperaban 3")
    if removed["Puntos"] != len(point_ids):
        raise ValueError(
            f"Se eliminaron {removed['Puntos']} puntos; se esperaban {len(point_ids)}"
        )

    retained = Counter()
    for sheet_name in ("Poligonos", "Puntos"):
        worksheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        for row_number in range(2, worksheet.max_row + 1):
            if worksheet.cell(row_number, headers["id"]).value:
                retained[worksheet.cell(row_number, headers["sector_investigado"]).value] += 1

    update_summary(workbook, retained, len(point_ids))

    methodology = workbook["Fuentes_y_metodologia"]
    methodology.append([])
    methodology.append(
        [
            "DEPURACION_ESPACIAL",
            "Eliminación de representaciones superpuestas",
            "",
            "Se retiraron tres edificios internos y los puntos cubiertos por polígonos conservados.",
        ]
    )
    workbook.save(OUTPUT_XLSX)

    print(f"XLSX: {OUTPUT_XLSX}")
    print(f"Edificios eliminados: {removed['Poligonos']}")
    print(f"Puntos eliminados: {removed['Puntos']}")
    print(f"Registros conservados: {sum(retained.values())}")
    for point_id in sorted(point_hits):
        print(f"  {point_id} | {point_names[point_id]} | {', '.join(point_hits[point_id])}")


if __name__ == "__main__":
    main()
