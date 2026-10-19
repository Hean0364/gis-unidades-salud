"""Genera la capa del mapa desde el XLSX sin unidades públicas."""

from __future__ import annotations

import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
EXCEL = (
    DATA
    / "unidades_salud_limpias_excel"
    / "unidades_salud_sin_publicas_sin_superposiciones.xlsx"
)
MAP_DIR = DATA / "mapa_san_salvador"
GEOMETRY_SOURCE = MAP_DIR / "unidades_salud_deduplicadas.geojson"
OUTPUT_GEOJSON = MAP_DIR / "unidades_salud_sin_publicas_sin_superposiciones.geojson"
OUTPUT_JS = MAP_DIR / "unidades_salud_sin_publicas_sin_superposiciones.js"
PUBLIC_SECTOR = "Pública"


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def read_rows() -> list[tuple[str, dict[str, Any]]]:
    workbook = load_workbook(EXCEL, read_only=True, data_only=True)
    records: list[tuple[str, dict[str, Any]]] = []

    for sheet_name in ("Poligonos", "Puntos"):
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        for values in rows:
            record = {
                header: json_value(value)
                for header, value in zip(headers, values)
                if value is not None and header != "geometria_wkt"
            }
            if not record.get("id"):
                continue
            if record.get("sector_investigado") == PUBLIC_SECTOR:
                raise ValueError(
                    f"El XLSX todavía contiene una unidad pública: {record['id']}"
                )
            records.append((sheet_name, record))

    workbook.close()
    return records


def map_properties(record: dict[str, Any], sheet_name: str) -> dict[str, Any]:
    properties = dict(record)
    original_name = properties.get("nombre")
    original_type = properties.get("tipo")
    original_sector = properties.get("sector")
    sector = properties.get("sector_investigado") or original_sector or "No determinada"

    properties["nombre_original"] = original_name
    properties["tipo_original"] = original_type
    properties["sector_original"] = original_sector
    properties["nombre"] = properties.get("nombre_rectificado") or original_name or "Sin nombre"
    properties["tipo"] = properties.get("tipo_rectificado") or original_type or "Sin tipo"
    properties["sector"] = sector
    properties["fuente"] = "Unificado filtrado"
    properties["representacion"] = "Polígono" if sheet_name == "Poligonos" else "Punto"

    if sector == "Privada con fines de lucro":
        properties["es_privado"] = "Sí, con fines de lucro"
    elif sector == "Privada sin fines de lucro":
        properties["es_privado"] = "Sí, sin fines de lucro"
    else:
        properties["es_privado"] = "No determinado"

    return properties


def build_collection() -> dict[str, Any]:
    geometry_collection = json.loads(GEOMETRY_SOURCE.read_text(encoding="utf-8"))
    geometry_by_id = {
        feature.get("properties", {}).get("id"): feature.get("geometry")
        for feature in geometry_collection.get("features", [])
    }

    records = read_rows()
    features: list[dict[str, Any]] = []
    counts: Counter[str] = Counter()
    sectors: Counter[str] = Counter()
    missing_geometry: list[str] = []

    for sheet_name, record in records:
        record_id = str(record.get("id") or "")
        geometry = geometry_by_id.get(record_id)
        if not geometry:
            missing_geometry.append(record_id)
            continue

        geometry_type = geometry.get("type")
        expected = (
            {"Polygon", "MultiPolygon"}
            if sheet_name == "Poligonos"
            else {"Point", "MultiPoint"}
        )
        if geometry_type not in expected:
            raise ValueError(
                f"Geometría inesperada para {record_id}: "
                f"hoja {sheet_name}, tipo {geometry_type}"
            )

        counts[sheet_name] += 1
        sectors[str(record.get("sector_investigado") or "No determinada")] += 1
        features.append(
            {
                "type": "Feature",
                "properties": map_properties(record, sheet_name),
                "geometry": geometry,
            }
        )

    if missing_geometry:
        raise ValueError(
            f"Falta geometría para {len(missing_geometry)} registros: {missing_geometry}"
        )
    if len(features) != len(records):
        raise ValueError(f"Se esperaban {len(records)} registros y se generaron {len(features)}")

    return {
        "type": "FeatureCollection",
        "name": "unidades_salud_san_salvador_sin_publicas_sin_superposiciones",
        "meta": {
            "departamento": "San Salvador",
            "fuente_tabular": str(EXCEL.relative_to(ROOT)).replace("\\", "/"),
            "fuente_geometria": str(GEOMETRY_SOURCE.relative_to(ROOT)).replace("\\", "/"),
            "filtros": [
                "Excluye registros clasificados como Pública; conserva las unidades privadas y No determinada",
                "Los tres edificios internos identificados eran públicos y quedaron excluidos por el filtro sectorial",
                "Excluye puntos ubicados dentro o sobre el borde de polígonos conservados",
            ],
            "edificios_internos_publicos_excluidos_por_sector": 3,
            "puntos_sobre_poligonos_eliminados": 2,
            "count": len(features),
            "poligonos": counts["Poligonos"],
            "puntos": counts["Puntos"],
            "sectores": dict(sectors),
            "capas_mapa": ["Polígonos", "Puntos"],
        },
        "features": features,
    }


def main() -> None:
    collection = build_collection()
    OUTPUT_GEOJSON.write_text(
        json.dumps(collection, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    OUTPUT_JS.write_text(
        "window.UNIDADES_DEPURADAS_GEOJSON = "
        + json.dumps(collection, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    meta = collection["meta"]
    print(f"GeoJSON: {OUTPUT_GEOJSON}")
    print(f"JavaScript: {OUTPUT_JS}")
    print(
        f"Registros: {meta['count']} "
        f"(polígonos: {meta['poligonos']}, puntos: {meta['puntos']})"
    )
    print(f"Sectores: {meta['sectores']}")


if __name__ == "__main__":
    main()
