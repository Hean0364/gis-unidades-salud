"""Genera 13 mapas independientes a partir de los Excel departamentales filtrados.

Ejecutar: python "Unidades de Salud/scripts/generar_mapas_departamentales.py"
Los GeoJSON se pueden usar en SIG; los JS permiten abrir index.html con file://.
No modifica los Excel ni el CSV maestro.
En La Paz y La Libertad se incluyen únicamente las filas verdes del Excel.
"""

from __future__ import annotations

import json
import math
import re
from collections import Counter
from html import escape
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DEPARTMENTS = DATA / "por_departamento"
BOUNDARIES = DATA / "geojson" / "departamentos.geojson"
REQUIRED = {"id", "nombre", "tipo", "departamento", "geometria_wkt", "centroide_latitud", "centroide_longitud"}
GREEN_ONLY_DEPARTMENTS = {"La Paz", "La Libertad"}
GREEN_RGB = {"00B050", "92D050", "0B6E4F"}


def coords(pair: str) -> list[float]:
    lon, lat = (float(value) for value in pair.strip().split())
    if not (math.isfinite(lon) and math.isfinite(lat) and -180 <= lon <= 180 and -90 <= lat <= 90):
        raise ValueError(f"Coordenadas inválidas: {pair}")
    return [lon, lat]


def geometry(wkt: str) -> dict:
    if wkt.startswith("POINT (") and wkt.endswith(")"):
        return {"type": "Point", "coordinates": coords(wkt[7:-1])}
    if wkt.startswith("POLYGON (") and wkt.endswith(")"):
        rings = re.findall(r"\(([^()]+)\)", wkt[9:-1])
        if not rings:
            raise ValueError(f"Polígono sin anillos: {wkt[:100]}")
        coordinates = [[coords(pair) for pair in ring.split(",")] for ring in rings]
        if any(len(ring) < 4 or ring[0] != ring[-1] for ring in coordinates):
            raise ValueError(f"Polígono sin anillo cerrado: {wkt[:100]}")
        return {"type": "Polygon", "coordinates": coordinates}
    raise ValueError(f"Geometría no admitida: {wkt[:100]}")


def is_green_row(row) -> bool:
    return any(
        cell.value is not None
        and cell.fill.fill_type == "solid"
        and cell.fill.fgColor.type == "rgb"
        and cell.fill.fgColor.rgb[-6:].upper() in GREEN_RGB
        for cell in row
    )


def records(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows()
        # Excel puede conservar formato en columnas vacías a la derecha de la tabla.
        columns = [(index, cell.value) for index, cell in enumerate(next(rows)) if cell.value is not None]
        names = [name for _, name in columns]
        if len(names) != len(set(names)) or not REQUIRED.issubset(names):
            raise ValueError(f"Columnas inesperadas en {path}")
        green_only = path.parent.name in GREEN_ONLY_DEPARTMENTS
        return [
            {name: row[index].value for index, name in columns}
            for row in rows
            if any(row[index].value is not None for index, _ in columns)
            and (not green_only or is_green_row(row))
        ]
    finally:
        workbook.close()


def collection(department: str, rows: list[dict]) -> dict:
    features = []
    seen = set()
    for row in rows:
        record_id = row["id"]
        if not record_id or record_id in seen or row["departamento"] != department:
            raise ValueError(f"ID o departamento inválido en {department}: {record_id}")
        seen.add(record_id)
        try:
            shape = geometry(row["geometria_wkt"])
            # Los centroides permiten localizar también los polígonos pequeños en el mapa.
            center = [float(row["centroide_longitud"]), float(row["centroide_latitud"])]
            if not all(math.isfinite(v) for v in center):
                raise ValueError("Centroide inválido")
            if shape["type"] == "Point" and any(abs(a - b) > 1e-5 for a, b in zip(shape["coordinates"], center)):
                raise ValueError("Punto y centroide no coinciden")
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{department}: geometría inválida en {record_id}") from exc
        properties = {key: value for key, value in row.items() if key != "geometria_wkt"}
        # Si se corrigió el centroide, no conservar una búsqueda con la ubicación anterior.
        if str(properties.get("google_maps") or "").startswith("https://www.google.com/maps/search/?api=1&query="):
            properties["google_maps"] = (
                f"https://www.google.com/maps/search/?api=1&query={row['centroide_latitud']},{row['centroide_longitud']}"
            )
        features.append({
            "type": "Feature",
            "properties": properties,
            "geometry": shape,
        })
    counts = Counter(feature["geometry"]["type"] for feature in features)
    return {
        "type": "FeatureCollection",
        "name": f"establecimientos_salud_{department}",
        "meta": {
            "departamento": department,
            "fuente": f"data/por_departamento/{department}/establecimientos_salud.xlsx",
            "count": len(features),
            "puntos": counts["Point"],
            "poligonos": counts["Polygon"],
        },
        "features": features,
    }


def html_page(department: str) -> str:
    name = escape(department)
    return f'''<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Unidades de Salud · {name}</title>
  <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
  <link rel="stylesheet" href="../../../css/styles.css">
</head>
<body>
  <aside class="dock">
    <section class="dock-section dock-layers">
      <header class="dock-title">
        <span>Capas · {name}</span>
        <div class="dock-tools">
          <button id="expand-all" type="button" title="Expandir todo" aria-label="Expandir todo">
            <svg viewBox="0 0 16 16"><rect x="2.5" y="2.5" width="11" height="11" fill="#fff" stroke="currentColor"/><path d="M5 8h6M8 5v6" stroke="currentColor" stroke-width="1.4"/></svg>
          </button>
          <button id="collapse-all" type="button" title="Contraer todo" aria-label="Contraer todo">
            <svg viewBox="0 0 16 16"><rect x="2.5" y="2.5" width="11" height="11" fill="#fff" stroke="currentColor"/><path d="M5 8h6" stroke="currentColor" stroke-width="1.4"/></svg>
          </button>
          <span class="tool-sep"></span>
          <button id="zoom-layer" type="button" title="Zoom a la capa seleccionada" aria-label="Zoom a la capa seleccionada">
            <svg viewBox="0 0 16 16"><circle cx="6.5" cy="6.5" r="4" fill="#fff" stroke="currentColor" stroke-width="1.4"/><path d="M9.5 9.5l4.5 4.5" stroke="currentColor" stroke-width="1.8"/><rect x="4.5" y="4.5" width="4" height="4" fill="#2f80ed"/></svg>
          </button>
          <button id="full-extent" type="button" title="Extensión completa" aria-label="Extensión completa">
            <svg viewBox="0 0 16 16"><path d="M2 6V2h4M10 2h4v4M14 10v4h-4M6 14H2v-4" fill="none" stroke="currentColor" stroke-width="1.4"/><rect x="5" y="5" width="6" height="6" fill="#7fa86f"/></svg>
          </button>
        </div>
      </header>
      <div class="map-search">
        <label class="visually-hidden" for="map-search-input">Buscar establecimiento</label>
        <div class="search-field">
          <input id="map-search-input" type="search" placeholder="Buscar establecimiento…" autocomplete="off"
            role="combobox" aria-autocomplete="list" aria-expanded="false" aria-controls="map-search-results">
          <button id="map-search-clear" type="button" aria-label="Limpiar búsqueda" title="Limpiar búsqueda">×</button>
        </div>
        <div id="map-search-results" class="search-results" role="listbox" aria-label="Resultados de búsqueda" hidden></div>
      </div>
      <div class="tree" id="layers" role="tree"></div>
    </section>
  </aside>
  <div id="map"></div>
  <footer class="statusbar">
    <span id="status">Cargando capas…</span>
    <span class="statusbar-right"><span id="coords">Lat —  Lon —</span><span class="crs">EPSG:4326</span></span>
  </footer>
  <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
  <script src="../../../js/mapas_base.js"></script>
  <script src="establecimientos.js"></script>
  <script src="limite_departamento.js"></script>
  <script src="../../../js/app_departamento.js"></script>
</body>
</html>
'''


def write_layer(path: Path, variable: str, data: dict) -> None:
    text = json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    path.with_suffix(".geojson").write_text(text, encoding="utf-8")
    path.with_suffix(".js").write_text(
        f"window.{variable} = " + json.dumps(data, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + ";\n",
        encoding="utf-8",
    )


def main() -> None:
    boundaries = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    by_name = {feature["properties"]["nombre"]: feature for feature in boundaries["features"]}
    paths = sorted(DEPARTMENTS.glob("*/establecimientos_salud.xlsx"))
    if len(paths) != 13 or {p.parent.name for p in paths} != set(by_name) - {"San Salvador"}:
        raise ValueError("Los Excel no corresponden a los 13 departamentos esperados")
    total = 0
    for excel in paths:
        department = excel.parent.name
        data = collection(department, records(excel))
        boundary = {
            "type": "FeatureCollection",
            "name": f"limite_{department}",
            "meta": {"fuente": boundaries.get("meta"), "departamento": department, "count": 1},
            "features": [by_name[department]],
        }
        write_layer(excel.parent / "establecimientos", "ESTABLECIMIENTOS_DEPARTAMENTO_GEOJSON", data)
        write_layer(excel.parent / "limite_departamento", "LIMITE_DEPARTAMENTO_GEOJSON", boundary)
        (excel.parent / "index.html").write_text(html_page(department), encoding="utf-8")
        total += data["meta"]["count"]
        print(f"{department}: {data['meta']['count']} registros, 1 límite")
    print(f"Total: {total} registros en {len(paths)} mapas")


if __name__ == "__main__":
    main()
