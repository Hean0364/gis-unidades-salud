"""Exporta únicamente los polígonos verdes del XLSX y todos los del CSV."""

from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parent
SOURCE_XLSX = WORKSPACE / "unidades_salud_sin_publicas_sin_superposiciones_2.xlsx"
SOURCE_CSV = WORKSPACE / "poligonos heber.csv"
OUTPUT_XLSX = ROOT / "data" / "unidades_salud_limpias_excel" / "unidades_salud_privadas_finales.xlsx"
OUTPUT_GEOJSON = ROOT / "data" / "mapa_san_salvador" / "privadas_finales.geojson"
OUTPUT_JS = ROOT / "data" / "mapa_san_salvador" / "privadas_finales.js"

GREEN_RGB = {"00B050", "92D050", "0B6E4F"}
EXTRA_COLUMNS = ["seleccion_xlsx", "fuente_geometria"]
COORD_RE = re.compile(r"(-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)\s+(-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)")


def is_green_row(row: Any) -> bool:
    for cell in row:
        if cell.value is None or cell.fill.fill_type != "solid":
            continue
        color = cell.fill.fgColor
        if color.type == "rgb" and color.rgb and color.rgb[-6:].upper() in GREEN_RGB:
            return True
    return False


def split_top_level(text: str) -> list[str]:
    parts: list[str] = []
    depth = 0
    start = 0
    for index, char in enumerate(text):
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif char == "," and depth == 0:
            parts.append(text[start:index].strip())
            start = index + 1
    parts.append(text[start:].strip())
    return [part for part in parts if part]


def strip_outer(text: str) -> str:
    text = text.strip()
    if not text.startswith("(") or not text.endswith(")"):
        raise ValueError(f"Paréntesis WKT inválidos: {text[:100]}")
    depth = 0
    for index, char in enumerate(text):
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth == 0 and index != len(text) - 1:
                raise ValueError(f"Paréntesis WKT inválidos: {text[:100]}")
    return text[1:-1].strip()


def parse_ring(text: str) -> list[list[float]]:
    coordinates = [[float(x), float(y)] for x, y in COORD_RE.findall(strip_outer(text))]
    if len(coordinates) < 4:
        raise ValueError(f"Anillo WKT incompleto: {text[:100]}")
    return coordinates


def parse_polygon_body(text: str) -> list[list[list[float]]]:
    return [parse_ring(ring) for ring in split_top_level(text)]


def parse_polygon_wkt(wkt: str) -> dict[str, Any]:
    value = str(wkt or "").strip()
    match = re.match(r"^(POLYGON|MULTIPOLYGON)\s*\(", value, flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"Se esperaba un polígono WKT: {value[:100]}")
    body = value[match.end() : -1].strip()
    if match.group(1).upper() == "POLYGON":
        return {"type": "Polygon", "coordinates": parse_polygon_body(body)}
    polygons = [parse_polygon_body(strip_outer(item)) for item in split_top_level(body)]
    return {"type": "MultiPolygon", "coordinates": polygons}


def ring_area_centroid(ring: list[list[float]]) -> tuple[float, float, float]:
    cross_sum = cx_sum = cy_sum = 0.0
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]):
        cross = x1 * y2 - x2 * y1
        cross_sum += cross
        cx_sum += (x1 + x2) * cross
        cy_sum += (y1 + y2) * cross
    if abs(cross_sum) < 1e-18:
        return 0.0, 0.0, 0.0
    return cx_sum / (3 * cross_sum), cy_sum / (3 * cross_sum), abs(cross_sum / 2)


def polygon_centroid(rings: list[list[list[float]]]) -> tuple[float, float, float]:
    total_area = weighted_x = weighted_y = 0.0
    for index, ring in enumerate(rings):
        x, y, area = ring_area_centroid(ring)
        signed_area = area if index == 0 else -area
        total_area += signed_area
        weighted_x += x * signed_area
        weighted_y += y * signed_area
    if total_area <= 0:
        ring = rings[0]
        return sum(p[0] for p in ring) / len(ring), sum(p[1] for p in ring) / len(ring), 0.0
    return weighted_x / total_area, weighted_y / total_area, total_area


def geometry_centroid(geometry: dict[str, Any]) -> tuple[float, float]:
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    parts = [polygon_centroid(rings) for rings in polygons]
    area_total = sum(area for _, _, area in parts)
    if area_total <= 0:
        return parts[0][0], parts[0][1]
    return (
        sum(x * area for x, _, area in parts) / area_total,
        sum(y * area for _, y, area in parts) / area_total,
    )


def geometry_vertex_count(geometry: dict[str, Any]) -> int:
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    return sum(len(ring) for polygon in polygons for ring in polygon)


def read_green_polygons() -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(SOURCE_XLSX, data_only=True)
    sheet = workbook["Poligonos"]
    headers = [str(cell.value) for cell in sheet[1]]
    records: list[dict[str, Any]] = []
    for row_number in range(2, sheet.max_row + 1):
        cells = sheet[row_number]
        record = {header: cell.value for header, cell in zip(headers, cells)}
        if not record.get("id") or not is_green_row(cells):
            continue
        parse_polygon_wkt(str(record.get("geometria_wkt") or ""))
        record["seleccion_xlsx"] = "Verde"
        record["fuente_geometria"] = "XLSX · hoja Poligonos"
        records.append(record)
    workbook.close()
    return headers, records


def csv_polygon_records(headers: list[str]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []
    with SOURCE_CSV.open("r", encoding="utf-8-sig", newline="") as source:
        for line_number, row in enumerate(csv.reader(source), start=1):
            if not row or not row[0].strip().upper().startswith(("POLYGON", "MULTIPOLYGON")):
                continue
            wkt = row[0].strip()
            geometry = parse_polygon_wkt(wkt)
            name = row[1].strip() if len(row) > 1 else f"Polígono CSV {line_number}"
            description = ", ".join(part.strip() for part in row[2:] if part.strip())
            centroid_lon, centroid_lat = geometry_centroid(geometry)
            record = {header: None for header in headers}
            record.update(
                {
                    "id": f"csv-heber-poligono-{line_number}",
                    "fuente": "CSV",
                    "fuente_id": f"poligonos heber.csv:{line_number}",
                    "fuente_conservada": "CSV · poligonos heber.csv",
                    "nombre": name,
                    "tipo": "Polígono de establecimiento de salud",
                    "tipo_general": "Salud",
                    "detalle": "Polígono",
                    "detalle_codigo": "poligono",
                    "departamento": "San Salvador",
                    "sector": "No determinada",
                    "sector_investigado": "No determinada",
                    "geometria_tipo": geometry["type"],
                    "vertices": geometry_vertex_count(geometry),
                    "centroide_latitud": centroid_lat,
                    "centroide_longitud": centroid_lon,
                    "centroide_metodo": "centroide_area",
                    "geometria_wkt": wkt,
                    "descripcion_breve": description or None,
                    "seleccion_xlsx": "CSV",
                    "fuente_geometria": f"CSV · {SOURCE_CSV.name}",
                }
            )
            records.append(record)
            audit.append(
                {
                    "fila_csv": line_number,
                    "nombre": name,
                    "descripcion": description,
                    "decision": "Incluido como polígono independiente",
                    "geometria_tipo": geometry["type"],
                    "WKT": wkt,
                }
            )
    return records, audit


def map_properties(record: dict[str, Any]) -> dict[str, Any]:
    properties = {key: value for key, value in record.items() if value is not None and key != "geometria_wkt"}
    original_name = properties.get("nombre")
    original_type = properties.get("tipo")
    properties["nombre_original"] = original_name
    properties["tipo_original"] = original_type
    properties["nombre"] = properties.get("nombre_rectificado") or original_name or "Sin nombre"
    properties["tipo"] = properties.get("tipo_rectificado") or original_type or "Sin tipo"
    properties["sector"] = properties.get("sector_investigado") or "No determinada"
    properties["representacion"] = "Polígono"
    properties["es_privado"] = (
        "Sí, con fines de lucro"
        if properties["sector"] == "Privada con fines de lucro"
        else "Sí, sin fines de lucro"
        if properties["sector"] == "Privada sin fines de lucro"
        else "No determinado"
    )
    return properties


def make_feature(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "type": "Feature",
        "properties": map_properties(record),
        "geometry": parse_polygon_wkt(str(record["geometria_wkt"])),
    }


def style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, cell in enumerate(sheet[1], start=1):
        header = str(cell.value or "")
        width = min(max(len(header) + 2, 14), 30)
        if header == "geometria_wkt":
            width = 72
        elif header in {"descripcion_breve", "evidencia_rectificacion", "criterio_sector"}:
            width = 48
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_data_sheet(workbook: Workbook, title: str, headers: list[str], records: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(header) for header in headers])
    style_sheet(sheet)


def write_workbook(headers: list[str], polygons: list[dict[str, Any]], audit: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    rows = [
        ["Polígonos finales · XLSX verde + CSV"],
        ["Fuente XLSX", SOURCE_XLSX.name],
        ["Fuente CSV", SOURCE_CSV.name],
        ["Polígonos verdes del XLSX", sum(r.get("fuente_geometria") == "XLSX · hoja Poligonos" for r in polygons)],
        ["Polígonos del CSV", sum(r.get("fuente_geometria") == f"CSV · {SOURCE_CSV.name}" for r in polygons)],
        ["Total de polígonos", len(polygons)],
        ["Puntos incluidos", 0],
        ["Criterio", "Se incluyen todas las filas verdes de la hoja Poligonos del XLSX y todos los polígonos del CSV. No se incorporan puntos."],
        ["Geometrías", "Los polígonos existentes del XLSX se conservan; los CSV se agregan por separado, sin asignarlos ni fusionarlos con puntos."],
    ]
    for row in rows:
        summary.append(row)
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    summary["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    summary["A1"].fill = PatternFill("solid", fgColor="0F766E")
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 28
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 110
    for row in summary.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color="1F2937")
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    summary.freeze_panes = "A2"

    write_data_sheet(workbook, "Poligonos", headers + EXTRA_COLUMNS, polygons)
    write_data_sheet(
        workbook,
        "Revision_CSV",
        ["fila_csv", "nombre", "descripcion", "decision", "geometria_tipo", "WKT"],
        audit,
    )
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_XLSX)


def build_outputs() -> dict[str, Any]:
    if not SOURCE_XLSX.exists():
        raise FileNotFoundError(SOURCE_XLSX)
    if not SOURCE_CSV.exists():
        raise FileNotFoundError(SOURCE_CSV)
    headers, xlsx_polygons = read_green_polygons()
    csv_records, audit = csv_polygon_records(headers)
    polygons = xlsx_polygons + csv_records
    features = [make_feature(record) for record in polygons]
    collection = {
        "type": "FeatureCollection",
        "name": "privadas_finales_poligonos",
        "meta": {
            "fuente_xlsx": SOURCE_XLSX.name,
            "fuente_csv": SOURCE_CSV.name,
            "poligonos_xlsx_verdes": len(xlsx_polygons),
            "poligonos_csv": len(csv_records),
            "poligonos": len(features),
            "puntos": 0,
            "count": len(features),
        },
        "features": features,
    }
    OUTPUT_GEOJSON.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_GEOJSON.write_text(json.dumps(collection, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUTPUT_JS.write_text(
        "window.PRIVADAS_GEOJSON = "
        + json.dumps(collection, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    write_workbook(headers, polygons, audit)
    return collection


def main() -> None:
    collection = build_outputs()
    meta = collection["meta"]
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"GeoJSON: {OUTPUT_GEOJSON}")
    print(f"JavaScript: {OUTPUT_JS}")
    print(
        f"Capa: {meta['poligonos']} polígonos "
        f"({meta['poligonos_xlsx_verdes']} XLSX verdes + {meta['poligonos_csv']} CSV); 0 puntos"
    )


if __name__ == "__main__":
    main()
