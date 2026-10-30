"""Elimina duplicados confirmados de un Excel departamental y regenera su capa.

Conserva sedes distintas con nombres iguales y registra las filas retiradas.
Solo modifica el departamento recibido; no ejecuta tests.
"""

from collections import Counter, defaultdict
from datetime import datetime
import json
import os
from pathlib import Path
import re
import shutil
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from generar_mapas_departamentales import DEPARTMENTS, collection, geometry, is_green_row, records, write_layer


def normalized_name(value):
    value = unicodedata.normalize("NFD", str(value or "").casefold())
    return " ".join("".join(char for char in value if not unicodedata.combining(char)).split())


def ring_key(ring):
    points = tuple(tuple(pair) for pair in ring)
    if points[0] == points[-1]:
        points = points[:-1]
    return min(sequence[index:] + sequence[:index] for sequence in (points, points[::-1]) for index in range(len(points)))


def geometry_key(shape):
    if shape["type"] == "Point":
        return "Point", tuple(shape["coordinates"])
    rings = shape["coordinates"]
    return "Polygon", ring_key(rings[0]), tuple(sorted(ring_key(ring) for ring in rings[1:]))


def deduplicate_department(department):
    folder = DEPARTMENTS / department
    path = folder / "establecimientos_salud.xlsx"
    audit_path = folder / "duplicados_eliminados.json"
    workbook = load_workbook(path)
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    entries = []
    for cells in sheet.iter_rows(min_row=2):
        record = {name: cells[column - 1].value for name, column in columns.items()}
        if record.get("id"):
            comment = cells[columns["geometria_wkt"] - 1].comment
            manual_geometry = record.get("fuente") == "manual" or bool(comment and "usuario" in normalized_name(comment.text))
            entries.append({"row": cells[0].row, "record": record, "green": is_green_row(cells), "manual_geometry": manual_geometry})
    parents = list(range(len(entries)))

    def root(index):
        while parents[index] != index:
            parents[index] = parents[parents[index]]
            index = parents[index]
        return index

    source_ids = {}
    same_places = {}
    generic_names = {"", "sin nombre", "clinica medica", "unidad de salud", "isbm"}
    for index, entry in enumerate(entries):
        record = entry["record"]
        source_id = str(record.get("fuente_id") or "")
        if re.fullmatch(r"(?:node|way|relation)/\d+", source_id) and record.get("fuente") in {"overpass", "geofabrik", "hot"}:
            if source_id in source_ids:
                parents[root(index)] = root(source_ids[source_id])
            source_ids[source_id] = index
        name = normalized_name(record.get("nombre"))
        if name not in generic_names:
            key = name, geometry_key(geometry(record["geometria_wkt"]))
            if key in same_places:
                parents[root(index)] = root(same_places[key])
            same_places[key] = index
    groups = defaultdict(list)
    for index, entry in enumerate(entries):
        groups[root(index)].append(entry)

    priority = {"manual": 5, "overpass": 4, "minsal": 3, "hot": 2, "geofabrik": 1}

    def rank(entry):
        record = entry["record"]
        return (
            entry["manual_geometry"],
            record.get("geometria_tipo") == "Polygon",
            entry["green"],
            sum(value not in (None, "") for value in record.values()),
            priority.get(record.get("fuente"), 0),
            -entry["row"],
        )

    protected = {"id", "fuente", "fuente_id", "geometria_wkt", "geometria_tipo", "vertices", "centroide_latitud", "centroide_longitud", "centroide_metodo", "google_maps", "detalle", "detalle_codigo"}
    removed = []
    for group in groups.values():
        if len(group) < 2:
            continue
        group.sort(key=rank, reverse=True)
        keeper = group[0]
        kept = dict(keeper["record"])
        for duplicate in group[1:]:
            for name, value in duplicate["record"].items():
                if name not in protected and kept.get(name) in (None, "") and value not in (None, ""):
                    kept[name] = value
                    sheet.cell(keeper["row"], columns[name], value)
            removed.append({
                "id_conservado": kept["id"],
                "fila_original": duplicate["row"],
                "motivo": "Mismo objeto OSM" if kept.get("fuente_id") == duplicate["record"].get("fuente_id") else "Mismo nombre específico y misma geometría exacta",
                "registro_eliminado": duplicate["record"],
            })
    if removed:
        slug = normalized_name(department).replace(" ", "-")
        backup = Path(os.environ["LOCALAPPDATA"]) / "Temp" / "opencode" / f"{slug}-antes-deduplicar-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
        shutil.copy2(path, backup)
        audit = json.loads(audit_path.read_text(encoding="utf-8")) if audit_path.exists() else {"departamento": department, "registros_eliminados": []}
        audit["registros_eliminados"].extend(removed)
        audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        for row_number in sorted({entry["fila_original"] for entry in removed}, reverse=True):
            sheet.delete_rows(row_number)
        extent = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        sheet.auto_filter.ref = extent
        for table in sheet.tables.values():
            table.ref = extent
            if table.autoFilter is not None:
                table.autoFilter.ref = extent
        workbook.save(path)
        print(f"Copia de seguridad: {backup}")
    workbook.close()
    data = collection(department, records(path))
    write_layer(folder / "establecimientos", "ESTABLECIMIENTOS_DEPARTAMENTO_GEOJSON", data)
    sectors = Counter(feature["properties"].get("sector", "No determinado") for feature in data["features"])
    print(f"Duplicados eliminados: {len(removed)}")
    print(f"{department}: {data['meta']['count']} registros; {data['meta']['puntos']} puntos; {data['meta']['poligonos']} polígonos")
    print(f"Sectores: {dict(sectors)}")
    for entry in removed:
        print(f"{entry['registro_eliminado']['id']} -> {entry['id_conservado']}")
    return {
        "departamento": department,
        "antes": len(entries),
        "eliminados": len(removed),
        "despues": data["meta"]["count"],
        "collection": data,
    }

