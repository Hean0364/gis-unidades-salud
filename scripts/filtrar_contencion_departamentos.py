"""Retira de los Excel departamentales puntos y polígonos dentro de polígonos mayores.

Requiere: pip install openpyxl "shapely>=2"
Ejecutar: python "Unidades de Salud/scripts/filtrar_contencion_departamentos.py"

Solo se compara dentro del mismo departamento. Una intersección parcial, un
polígono del mismo tamaño o un punto sobre el borde NO se consideran contenido.
Conserva el registro completo de cada fila retirada en un CSV de auditoría.
El CSV maestro, los límites y los archivos HTML/JS de los visores no se alteran.
"""

from __future__ import annotations

import csv
import hashlib
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from shapely.geometry import shape
from shapely.strtree import STRtree

from generar_mapas_departamentales import DEPARTMENTS, ROOT, collection, geometry, records, write_layer


AUDIT = DEPARTMENTS / "filtrados_por_contencion.csv"
MASTER = ROOT / "data" / "normalizados" / "establecimientos_salud.csv"
MORE_COLUMNS = (
    "poligono_contenedor_id",
    "poligono_contenedor_nombre",
    "numero_contenedores",
    "criterio_contencion",
)


@dataclass
class Entry:
    path: Path
    row: dict
    geom: object


def load_entries() -> tuple[list[Path], dict[Path, list[dict]], list[Entry]]:
    paths = sorted(DEPARTMENTS.glob("*/establecimientos_salud.xlsx"))
    if len(paths) != 13:
        raise ValueError(f"Se esperaban 13 Excel, se encontraron {len(paths)}")
    rows_by_path = {path: records(path) for path in paths}
    entries = []
    ids = set()
    for path, rows in rows_by_path.items():
        for row in rows:
            record_id = row["id"]
            if not record_id or record_id in ids or row["departamento"] != path.parent.name:
                raise ValueError(f"ID o departamento inconsistente: {path} {record_id}")
            ids.add(record_id)
            geom = shape(geometry(row["geometria_wkt"]))
            if not geom.is_valid or geom.is_empty or geom.geom_type not in {"Polygon", "Point"}:
                raise ValueError(f"Geometría inválida en {path}: {record_id}")
            entries.append(Entry(path, row, geom))
    return paths, rows_by_path, entries


def contained(entries: list[Entry]) -> dict[str, list[Entry]]:
    """Registros contenidos y polígonos que los contienen (Shapely/GEOS)."""
    by_department: dict[str, list[Entry]] = {}
    for entry in entries:
        by_department.setdefault(entry.path.parent.name, []).append(entry)
    hits = {}
    for department, subset in by_department.items():
        parents = [entry for entry in subset if entry.geom.geom_type == "Polygon"]
        tree = STRtree([entry.geom for entry in parents])
        for entry in subset:
            matches = []
            for index in tree.query(entry.geom):
                parent = parents[int(index)]
                if parent is entry:
                    continue
                if entry.geom.geom_type == "Polygon" and parent.geom.area <= entry.geom.area + 1e-12:
                    continue
                if parent.geom.contains(entry.geom):
                    matches.append(parent)
            if matches:
                hits[entry.row["id"]] = matches
    return hits


def save_audit(entries: list[Entry], matches: dict[str, list[Entry]]) -> None:
    if AUDIT.exists():
        raise FileExistsError(f"No se sobrescribe un archivo de auditoría existente: {AUDIT}")
    columns = list(entries[0].row)
    with AUDIT.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns + list(MORE_COLUMNS))
        writer.writeheader()
        for entry in entries:
            parents = matches.get(entry.row["id"])
            if not parents:
                continue
            biggest = sorted(parents, key=lambda parent: (-parent.geom.area, parent.row["id"]))[0]
            writer.writerow({
                **entry.row,
                "poligono_contenedor_id": biggest.row["id"],
                "poligono_contenedor_nombre": biggest.row["nombre"],
                "numero_contenedores": len(parents),
                "criterio_contencion": "Polígono contenido o punto estrictamente dentro; mismo departamento; sin cruce parcial",
            })


def update_excel(path: Path, original: list[dict], remove_ids: set[str]) -> list[dict]:
    expected = [row for row in original if row["id"] not in remove_ids]
    if len(expected) == len(original):
        return expected
    book = load_workbook(path)
    sheet = book.active
    if sheet.max_row != len(original) + 1 or [cell.value for cell in sheet[1]] != list(original[0]):
        book.close()
        raise ValueError(f"El libro cambió desde su lectura: {path}")
    for line, row in reversed(list(enumerate(original, start=2))):
        if row["id"] in remove_ids:
            sheet.delete_rows(line)
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
    book.close()
    if records(path) != expected:
        raise AssertionError(f"La verificación de las filas conservadas falló: {path}")
    return expected


def regenerate(rows_by_path: dict[Path, list[dict]]) -> None:
    features = []
    departments = []
    for path, rows in rows_by_path.items():
        data = collection(path.parent.name, rows)
        write_layer(path.parent / "establecimientos", "ESTABLECIMIENTOS_DEPARTAMENTO_GEOJSON", data)
        features.extend(data["features"])
        departments.append(path.parent.name)
    counts = Counter(feature["geometry"]["type"] for feature in features)
    combined = {
        "type": "FeatureCollection",
        "name": "establecimientos_salud_13_departamentos",
        "meta": {
            "fuente": "data/por_departamento/*/establecimientos_salud.xlsx",
            "departamentos": departments,
            "count": len(features),
            "puntos": counts["Point"],
            "poligonos": counts["Polygon"],
        },
        "features": features,
    }
    write_layer(DEPARTMENTS / "establecimientos_todos", "ESTABLECIMIENTOS_TODOS_GEOJSON", combined)


def main() -> None:
    master_hash = hashlib.sha256(MASTER.read_bytes()).hexdigest()
    paths, rows_by_path, entries = load_entries()
    matches = contained(entries)
    if not matches:
        print("No hay nuevos registros contenidos. No se modificó ningún archivo.")
        return
    print(f"Encontrados: {len(matches)} registros contenidos (de {len(entries)} totales)")
    print("Por geometría:", dict(Counter(e.geom.geom_type for e in entries if e.row["id"] in matches)))
    print("Por departamento:", dict(Counter(e.path.parent.name for e in entries if e.row["id"] in matches)))
    save_audit(entries, matches)
    removed = set(matches)
    for path in paths:
        before = rows_by_path[path]
        rows_by_path[path] = update_excel(path, before, removed)
        print(f"{path.parent.name}: {len(before) - len(rows_by_path[path])} retirados; {len(rows_by_path[path])} restantes")

    regenerate(rows_by_path)
    _, _, survivors = load_entries()
    if contained(survivors):
        raise AssertionError("Aún quedan geometrías totalmente contenidas después del filtrado")
    if len(survivors) != len(entries) - len(matches):
        raise AssertionError("El total final no coincide con los registros retirados")
    if hashlib.sha256(MASTER.read_bytes()).hexdigest() != master_hash:
        raise AssertionError("El CSV maestro fue alterado")
    print(f"Total final: {len(survivors)}. Auditoría: {AUDIT}")


if __name__ == "__main__":
    main()
